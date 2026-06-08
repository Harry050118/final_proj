from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "个人贡献报告_郭锦添.docx"
SUMMARY = ROOT / "runs" / "clean_search" / "summary"
ASSETS = ROOT / "report_assets"


def set_east_asia_font(run, font_name="宋体"):
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False, color=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(str(text))
    set_east_asia_font(run)
    run.font.size = Pt(9)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def set_cell_margins(table, top=80, start=120, bottom=80, end=120):
    tbl_pr = table._tbl.tblPr
    tbl_cell_mar = tbl_pr.first_child_found_in("w:tblCellMar")
    if tbl_cell_mar is None:
        tbl_cell_mar = OxmlElement("w:tblCellMar")
        tbl_pr.append(tbl_cell_mar)
    for m, v in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tbl_cell_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tbl_cell_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def add_paragraph(doc, text="", style=None, bold_prefix=None):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_east_asia_font(r1)
        r1.bold = True
        r2 = p.add_run(text[len(bold_prefix):])
        set_east_asia_font(r2)
    else:
        r = p.add_run(text)
        set_east_asia_font(r)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(text)
    set_east_asia_font(r)
    return p


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(text)
    set_east_asia_font(r, "宋体")
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(89, 89, 89)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        set_east_asia_font(run, "黑体")
        run.font.color.rgb = RGBColor(31, 78, 121) if level > 1 else RGBColor(46, 116, 181)
    return p


def add_key_value_table(doc):
    table = doc.add_table(rows=3, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    set_cell_margins(table)
    rows = [("姓名", "郭锦添"), ("学号", "12311003"), ("小组", "Group Two")]
    for row, (k, v) in zip(table.rows, rows):
        set_cell_text(row.cells[0], k, bold=True)
        set_cell_text(row.cells[1], v)
        set_cell_shading(row.cells[0], "E8EEF5")
    for row in table.rows:
        row.cells[0].width = Cm(3.0)
        row.cells[1].width = Cm(9.0)
    return table


def add_metrics_table(doc):
    headers = ["实验设置", "Overall", "Content Coverage", "Citation Quality", "主要结论"]
    rows = [
        ["Baseline", "6.86", "4.09", "7.04", "能识别多数主题，但内容覆盖与引用可靠性不足。"],
        ["Skill Ablation", "7.70", "4.35", "8.98", "主要收益集中在引用质量与 claim-citation 检查。"],
        ["Kimi Ablation", "6.80", "3.91", "7.67", "局部提升引用质量，但整体表现不稳定。"],
    ]
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    set_cell_margins(table, top=100, bottom=100)
    for i, h in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], h, bold=True, color="FFFFFF")
        set_cell_shading(table.rows[0].cells[i], "1F4E79")
    for row_data in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row_data):
            set_cell_text(cells[i], val)
            if i == 0:
                set_cell_shading(cells[i], "E8EEF5")
    widths = [Cm(3.0), Cm(2.0), Cm(2.4), Cm(2.4), Cm(6.0)]
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = width
    return table


def add_image_if_exists(doc, path, width, caption):
    if path.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        run.add_picture(str(path), width=width)
        add_caption(doc, caption)
    else:
        add_paragraph(doc, f"（缺少图片文件：{path}）")


def configure_styles(doc):
    sec = doc.sections[0]
    sec.top_margin = Cm(2.2)
    sec.bottom_margin = Cm(2.0)
    sec.left_margin = Cm(2.3)
    sec.right_margin = Cm(2.3)

    normal = doc.styles["Normal"]
    normal.font.name = "宋体"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(6)

    for style_name in ["Heading 1", "Heading 2", "Heading 3"]:
        st = doc.styles[style_name]
        st.font.name = "黑体"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        st.font.bold = True
    doc.styles["Heading 1"].font.size = Pt(15)
    doc.styles["Heading 2"].font.size = Pt(13)
    doc.styles["Heading 3"].font.size = Pt(11.5)


def add_footer(doc):
    section = doc.sections[0]
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("个人贡献报告 - 郭锦添")
    set_east_asia_font(run)
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(100, 100, 100)


def build():
    doc = Document()
    configure_styles(doc)
    add_footer(doc)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(12)
    r = title.add_run("个人贡献报告")
    set_east_asia_font(r, "黑体")
    r.font.size = Pt(22)
    r.bold = True

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(18)
    r = subtitle.add_run("Group Two Final Project")
    set_east_asia_font(r, "Calibri")
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor(89, 89, 89)

    add_heading(doc, "一、基本信息", 1)
    add_key_value_table(doc)

    add_heading(doc, "二、个人贡献概述", 1)
    add_paragraph(
        doc,
        "本次课程项目中，我的工作重点集中在实验执行、实验结果核查、对比分析和展示材料整理。"
        "其中，我负责 Baseline 实验、Skill 消融实验和 Kimi 消融实验，并参与 Poster 的版式设计与视觉排版。"
    )
    add_paragraph(
        doc,
        "在实验侧，我需要保证不同设置之间的输入、评测流程和结果记录保持一致，使后续对比能够尽量反映变量本身的影响，"
        "而不是由样本选择、文件归档或评测口径差异造成。"
    )
    add_bullet(doc, "Baseline 实验：整理 OpenClaw + DS-v4-Pro 在无 Skill、无 Memory 条件下的五个样本结果，形成后续消融实验的对照基准。")
    add_bullet(doc, "Skill 消融实验：对比加入 citation-grounded workflow 后的表现，重点检查引用质量、claim-citation 对齐和错误引用减少情况。")
    add_bullet(doc, "Kimi 消融实验：在保持 OpenClaw 框架不变的情况下替换底层模型，观察模型选择对整体得分和引用可靠性的影响。")
    add_bullet(doc, "Poster 制作：参与版式与视觉排版，负责部分图表区域的摆放、层级组织和视觉一致性调整。")

    add_heading(doc, "三、实验工作与成果", 1)
    add_paragraph(
        doc,
        "下表汇总了我主要负责的三组实验。为了避免只看总体分数造成误判，我同时关注 Overall、Content Coverage 和 Citation Quality 三类指标。"
        "其中 Overall 反映综合表现，Content Coverage 反映生成内容对 gold related work 的覆盖程度，Citation Quality 则反映引用真实性、适切性、位置和主题一致性等问题。"
    )
    add_metrics_table(doc)
    add_caption(doc, "表 1  我负责的三组实验关键指标对比")

    add_heading(doc, "3.1 Baseline 实验", 2)
    add_paragraph(
        doc,
        "Baseline 采用 OpenClaw + DS-v4-Pro，不额外加入 Skill 或 Memory。我的主要工作是整理五个 paper-level task 的运行产物，"
        "核查候选 related work、候选 references、评测 JSON 与 Markdown 报告是否对应，并将结果汇总为后续所有消融实验的参照。"
    )
    add_paragraph(
        doc,
        "Baseline 的平均 Overall 为 6.86，Content Coverage 为 4.09，Citation Quality 为 7.04。这个结果说明模型能够识别多数论文主题，"
        "但仍然容易遗漏关键相关工作，也会出现 citation-claim mismatch 和 overclaim。典型问题包括：引用的论文只能弱支持周围论断、"
        "把通用推理或代码生成论文用于支持不完全相关的检索或 RLVR claim，以及在 related work 中遗漏关键研究脉络。"
    )

    add_heading(doc, "3.2 Skill 消融实验", 2)
    add_paragraph(
        doc,
        "Skill 消融实验的重点是验证 citation-grounded workflow 是否能改善引用可靠性。该设置要求在生成 related work 前进行候选论文筛选，"
        "并在输出前检查 citation-claim pair。我的工作包括执行 Skill 设置下的五个样本实验、整理评测报告，并将其与 Baseline 逐项对比。"
    )
    add_paragraph(
        doc,
        "实验结果显示，加入 Skill 后 Overall 从 6.86 提升到 7.70，Citation Quality 从 7.04 提升到 8.98。"
        "这说明 Skill 的主要价值不是让文本表面上更流畅，而是把 related work 生成转化为一个带约束的引用核查流程。"
        "不过，Skill 对 Content Coverage 的提升有限，仅从 4.09 提升到 4.35，说明它更擅长减少错误引用，而不是自动补全遗漏主题。"
    )

    add_heading(doc, "3.3 Kimi 消融实验", 2)
    add_paragraph(
        doc,
        "Kimi 消融实验用于观察仅替换底层模型是否能稳定提升表现。实验保持 OpenClaw 框架和任务设置不变，将模型替换为 Kimi-K2.6，"
        "从而尽量把变量集中在模型本身。我的工作包括运行五个样本、整理报告，并与 Baseline 对比 Overall、Citation Quality 和 Content Coverage。"
    )
    add_paragraph(
        doc,
        "结果显示，OpenClaw + Kimi-K2.6 的平均 Overall 为 6.80，略低于 Baseline 的 6.86；Citation Quality 为 7.67，高于 Baseline 的 7.04；"
        "但 Content Coverage 为 3.91，低于 Baseline 的 4.09。该结果说明更换模型可能带来局部指标改善，但不能稳定解决 related work 生成中的核心问题。"
        "尤其是在不同论文样本上，模型替换带来的收益并不一致，因此不能把可靠引用生成简单归因于模型能力。"
    )

    doc.add_page_break()
    add_heading(doc, "四、成果证据", 1)
    add_paragraph(
        doc,
        "以下图表来自项目实验汇总文件，用于证明三组实验的结果已经被纳入统一评测和展示。"
        "我在整理个人报告时保留了团队报告中的原始分数口径，避免重新计算造成口径不一致。"
    )
    add_image_if_exists(
        doc,
        SUMMARY / "clean_search_group_summary.png",
        Inches(5.7),
        "图 1  各实验组平均 Overall 得分对比",
    )
    add_image_if_exists(
        doc,
        SUMMARY / "clean_search_overall_by_task.png",
        Inches(5.7),
        "图 2  不同实验组在五个论文样本上的 Overall 得分",
    )

    add_heading(doc, "五、Poster 制作贡献", 1)
    add_paragraph(
        doc,
        "除实验部分外，我还参与了 Poster 的版式与视觉排版。我的工作重点不在实验内容撰写，而在于让已有实验结论、图表和文字说明在海报中形成清晰的信息层级。"
        "具体包括调整模块区域、图表摆放、标题层级、文字密度和视觉一致性，使读者能够先看到核心结论，再进一步阅读实验设置和结果解释。"
    )
    add_paragraph(
        doc,
        "在排版过程中，我重点关注三点：第一，保证结果图表与对应解释文本靠近，减少读者在版面中来回寻找信息；"
        "第二，控制不同区域的文字密度，避免海报变成报告正文的压缩版；第三，统一字体、颜色和图注风格，使海报整体更像一个完整展示材料。"
    )
    add_image_if_exists(
        doc,
        ASSETS / "poster_slide_1.png",
        Inches(6.0),
        "图 3  Poster 版式与视觉排版成果截图",
    )

    doc.add_page_break()
    add_heading(doc, "六、个人收获与反思", 1)
    add_heading(doc, "6.1 对消融实验和公平比较的理解", 2)
    add_paragraph(
        doc,
        "通过负责 Baseline、Skill 和 Kimi 三组实验，我更清楚地理解了消融实验的核心并不是简单多跑几个设置，而是要保证除目标变量外的其他条件尽量一致。"
        "Baseline 提供了判断改进是否有效的参照；Skill 实验用于验证工作流约束的作用；Kimi 实验则用于判断模型替换本身是否足够。"
        "如果没有统一的样本、统一的评测流程和统一的结果归档方式，后续结论就很容易失去可比性。"
    )
    add_heading(doc, "6.2 工程执行与结果归档能力", 2)
    add_paragraph(
        doc,
        "这次实验让我意识到，实验贡献不仅是得到一个分数，还包括保证运行产物可追溯。每个实验目录中的 s_text、s_reference、g_text、g_reference、report.json 和 report.md 都需要对应清楚。"
        "只有当这些文件能够相互校验时，后续的图表、团队报告和口头展示才有可靠依据。这个过程提升了我在批量实验执行、结果核查和材料整理方面的工程习惯。"
    )
    add_heading(doc, "6.3 从指标到结论的分析能力", 2)
    add_paragraph(
        doc,
        "我也学到不能只根据 Overall 得分判断一个实验是否成功。Skill 的 Overall 提升明显，但更关键的证据是 Citation Quality 的提升；"
        "Kimi 的 Citation Quality 有一定改善，但 Overall 和 Content Coverage 并没有稳定超过 Baseline。"
        "因此，分析实验结果时需要同时看总体指标、子指标和具体错误案例，才能判断改进到底作用在哪里。"
    )
    add_heading(doc, "6.4 对局限性的认识", 2)
    add_paragraph(
        doc,
        "从结果看，Skill 能明显降低错误引用风险，但对内容覆盖的帮助有限；Kimi 替换模型后局部指标改善，但整体不稳定。"
        "这说明 related work 生成中的 citation reliability 问题不能只依赖模型本身解决，也不能只依赖写作提示解决。"
        "更合理的方向可能是结合明确的引用检查流程、可复用的历史错误经验和更强的检索验证机制。"
    )
    add_heading(doc, "6.5 总结", 2)
    add_paragraph(
        doc,
        "总体而言，我在本项目中的主要贡献是完成三组关键对照实验，并将实验结果整理为可以支持团队结论的证据。"
        "这些工作帮助团队明确了 Baseline 的问题、Skill 的有效性以及 Kimi 模型替换的局限。"
        "同时，我也通过 Poster 排版工作参与了最终展示材料的组织，使实验发现能够以更清晰的视觉形式呈现。"
        "这次经历让我在实验设计、工程复现、结果分析和学术展示方面都有了更系统的训练。"
    )

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
